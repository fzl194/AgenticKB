"""File Inspector（M3 WP5, SRS §C03 / §4.2 字段裁剪版）.

职责：对冻结输入的原始 bytes 做确定性画像——detected MIME、source
format、容器数（页/sheet/slide）、加密标记、PDF 文本层抽样。产出
:class:`DocumentProfile` 供 Parser Router（§C05）与 Operator 审计使用。

设计要点：
- **复用** ``frozen_input/safe_intake`` 的签名探测（``SafeIntake.detect_mime``），
  不重复造轮子；仅在其返回 ``application/zip`` 时做最小扩展——读 ZIP 的
  ``[Content_Types].xml`` 消歧 OOXML 家族（docx/xlsx/pptx），使无
  declared_mime 的调用也能识别 Office 文档。
- 签名/文本启发**胜过** declared MIME（SRS §2.4 "扩展名不可信"）；未知
  格式输出 ``source_format="unknown"`` 而不是抛异常或偷偷当 TXT
  （SRS §4.4 unsupported 语义）。
- 全部从 ``io.BytesIO`` 打开（pdfplumber / python-docx / openpyxl /
  python-pptx），**不落盘**（§4.6：Adapter/Inspector 不触碰文件系统）。
- 纯同步函数，任何单探测失败降级为 warning，绝不让 Inspector 崩掉
  整条 intake 链路。
"""
from __future__ import annotations

import io
import zipfile
from dataclasses import dataclass, field

from knowledge_mining.mining.frozen_input.safe_intake import SafeIntake

#: 本 Inspector 的结构画像版本（画像语义变化时应递增，触发审计追溯）。
INSPECTOR_VERSION = "0.1.0"

#: declared MIME -> 合成文件名扩展提示（SafeIntake.detect_mime 的 hint 入参；
#: 签名命中时 hint 不参与判定，仅在无签名的文本格式间区分子类型）。
_MIME_EXT_HINT: dict[str, str] = {
    "text/plain": ".txt",
    "text/markdown": ".md",
    "text/html": ".html",
    "application/pdf": ".pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ".docx",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": ".xlsx",
    "application/vnd.openxmlformats-officedocument.presentationml.presentation": ".pptx",
}

#: detected MIME -> source_format 标签（与 SRS §C03 解析覆盖面对齐）。
_MIME_TO_FORMAT: dict[str, str] = {
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": "xlsx",
    "application/vnd.openxmlformats-officedocument.presentationml.presentation": "pptx",
    "text/html": "html",
    "text/markdown": "markdown",
    "text/plain": "txt",
}

#: [Content_Types].xml 中的 OOXML 命名空间标记 -> MIME（zip 消歧用）。
_OOXML_CONTENT_TYPE_MARKERS: tuple[tuple[str, str], ...] = (
    ("wordprocessingml", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
    ("spreadsheetml", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
    ("presentationml", "application/vnd.openxmlformats-officedocument.presentationml.presentation"),
)

#: PDF 文本层抽样的最大页数（SRS §4.2 "抽样前几页"）。
_PDF_TEXT_SAMPLE_PAGES = 3


@dataclass(frozen=True)
class DocumentProfile:
    """一次 inspect 的冻结画像（SRS §4.2 示例字段裁剪）.

    - ``container_count``：页/sheet/slide 数；取不到（docx 段落不是容器、
      文本格式、加密 PDF）为 ``None``——"未知可缺，不得伪造"（SRS §7.4）。
    - ``has_text_layer``：仅 PDF 抽样（前 3 页有无字符）；非 PDF 为 ``None``。
    - ``encrypted``：PDF 扫描到 ``/Encrypt`` 字典即标记；OOXML 家族在
      加密时实际退化为 OLE2 容器，会走 unknown 分支。
    - ``warnings``：降级/可疑事件（损坏 zip、打开失败等），供审计。
    """

    detected_mime: str
    source_format: str
    container_count: int | None = None
    encrypted: bool = False
    has_text_layer: bool | None = None
    container_kind: str | None = None
    inspector_version: str = INSPECTOR_VERSION
    warnings: tuple[str, ...] = field(default_factory=tuple)


class FileInspector:
    """无状态文档画像器（SRS §C03）."""

    __slots__ = ("_intake",)

    def __init__(self) -> None:
        # 复用 safe_intake 的签名探测（不重复实现魔法字节表）。
        self._intake = SafeIntake()

    def inspect(self, data: bytes, *, declared_mime: str | None = None) -> DocumentProfile:
        """对原始 ``data`` 产出 :class:`DocumentProfile`；永不因格式问题抛错."""
        mime = self._detect(data, declared_mime)
        if mime == "application/zip":
            mime = self._refine_ooxml(data)
        source_format = _MIME_TO_FORMAT.get(mime, "unknown")
        warnings: list[str] = []

        if source_format == "pdf":
            return self._profile_pdf(data, mime)
        if source_format in {"docx", "xlsx", "pptx"}:
            return self._profile_ooxml(data, mime, source_format, warnings)
        # 文本格式 / 未知格式：无容器语义
        if source_format == "unknown":
            warnings.append(f"unsupported mime for inspection: {mime}")
        return DocumentProfile(
            detected_mime=mime,
            source_format=source_format,
            warnings=tuple(warnings),
        )

    # -- MIME 探测 ----------------------------------------------------------

    def _detect(self, data: bytes, declared_mime: str | None) -> str:
        """SafeIntake 签名探测；declared MIME 只作扩展名提示（§2.4）."""
        ext = _MIME_EXT_HINT.get((declared_mime or "").lower(), "")
        filename = f"source{ext}"
        return self._intake.detect_mime(data, filename)

    @staticmethod
    def _refine_ooxml(data: bytes) -> str:
        """ZIP -> OOXML 家族消歧：读 ``[Content_Types].xml`` 标记.

        safe_intake 的 admission 层只用扩展名提示消歧（其注释已说明
        原因）；Inspector 拥有完整 bytes，故读一次内容类型表以支持无
        declared MIME 的调用。坏 zip 保持 ``application/zip`` -> unknown。
        """
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as zf:
                names = set(zf.namelist())
                if "[Content_Types].xml" not in names:
                    return "application/zip"
                content = zf.read("[Content_Types].xml")
        except (zipfile.BadZipFile, OSError, RuntimeError):
            return "application/zip"
        for marker, mime in _OOXML_CONTENT_TYPE_MARKERS:
            if marker.encode() in content:
                return mime
        return "application/zip"

    # -- PDF / OOXML 画像 ---------------------------------------------------

    def _profile_pdf(self, data: bytes, mime: str) -> DocumentProfile:
        """PDF：加密标记 + 页数 + 前 3 页字符抽样（SRS §4.2）."""
        encrypted = b"/Encrypt" in data
        if encrypted:
            return DocumentProfile(
                detected_mime=mime,
                source_format="pdf",
                encrypted=True,
                warnings=("pdf declares /Encrypt; parsing gated",),
            )
        import pdfplumber

        try:
            with pdfplumber.open(io.BytesIO(data)) as pdf:
                container_count = len(pdf.pages)
                sample = pdf.pages[:_PDF_TEXT_SAMPLE_PAGES]
                chars = sum(len(page.chars) for page in sample)
        except Exception as exc:  # pdfminer 异常族不稳定，统一降级
            return DocumentProfile(
                detected_mime=mime,
                source_format="pdf",
                warnings=(f"pdfplumber open failed: {type(exc).__name__}",),
            )
        return DocumentProfile(
            detected_mime=mime,
            source_format="pdf",
            container_count=container_count,
            encrypted=False,
            has_text_layer=chars > 0,
            container_kind="page",
        )

    def _profile_ooxml(
        self,
        data: bytes,
        mime: str,
        source_format: str,
        warnings: list[str],
    ) -> DocumentProfile:
        """OOXML：zipfile 容器校验 + 各库轻量打开取容器数（BytesIO）."""
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as zf:
                zf.testzip()  # None 即全部 CRC 通过
        except (zipfile.BadZipFile, OSError) as exc:
            warnings.append(f"invalid ooxml container: {type(exc).__name__}")
            return self._unknown(mime, warnings)

        try:
            container_count, container_kind = _ooxml_containers(data, source_format)
        except Exception as exc:
            warnings.append(
                f"{source_format} open failed: {type(exc).__name__}"
            )
            container_count, container_kind = None, None
        return DocumentProfile(
            detected_mime=mime,
            source_format=source_format,
            container_count=container_count,
            container_kind=container_kind,
            warnings=tuple(warnings),
        )

    @staticmethod
    def _unknown(mime: str, warnings: list[str]) -> DocumentProfile:
        return DocumentProfile(
            detected_mime=mime,
            source_format="unknown",
            warnings=tuple(warnings),
        )


def _ooxml_containers(data: bytes, source_format: str) -> tuple[int | None, str | None]:
    """按格式轻量打开 OOXML 取容器数；docx 段落不是容器返回 (None, None)."""
    buf = io.BytesIO(data)
    if source_format == "xlsx":
        import openpyxl

        # 本环境 openpyxl 的 read_only Workbook 不支持 with，显式 close
        workbook = openpyxl.load_workbook(buf, read_only=True)
        try:
            return len(workbook.sheetnames), "sheet"
        finally:
            workbook.close()
    if source_format == "pptx":
        from pptx import Presentation

        return len(Presentation(buf).slides), "slide"
    # docx：Document() 打开即校验可解析性；段落不构成原生容器（SRS §3.6）
    import docx

    docx.Document(buf)
    return None, None


__all__ = [
    "INSPECTOR_VERSION",
    "DocumentProfile",
    "FileInspector",
]
