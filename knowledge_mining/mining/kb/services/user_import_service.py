"""Small, synchronous CSV/XLSX intake for ordinary users and domain bindings."""
from __future__ import annotations

import csv
import io
import zipfile
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from psycopg.errors import UniqueViolation

from knowledge_mining.mining.infra.domain_pack import resolve_domain
from knowledge_mining.mining.kb.db import KbDB, UserDeletionConflict


MAX_IMPORT_BYTES = 5 * 1024 * 1024
MAX_IMPORT_ROWS = 5_000
MAX_XLSX_EXPANDED_BYTES = 50 * 1024 * 1024
MAX_XLSX_MEMBERS = 1_000


class UserImportError(ValueError):
    pass


class UserImportService:
    def __init__(self, db: KbDB) -> None:
        self._db = db

    async def execute(
        self,
        *,
        filename: str,
        content: bytes,
        dry_run: bool,
        domain: str | None = None,
    ) -> dict[str, Any]:
        rows = self._parse(filename=filename, content=content, domain_mode=domain is not None)
        plan = await self._plan(rows=rows, domain=domain)
        if not dry_run and not plan["errors"]:
            try:
                await self._db.apply_user_import(
                    new_users=plan["new_users"], bindings=plan["bindings"],
                )
            except (UniqueViolation, UserDeletionConflict) as exc:
                raise UserImportError("import_changed_since_preview") from exc
            plan = {**plan, "applied": True}
        else:
            plan = {**plan, "applied": False}
        return plan

    def _parse(
        self, *, filename: str, content: bytes, domain_mode: bool,
    ) -> list[dict[str, Any]]:
        if len(content) > MAX_IMPORT_BYTES:
            raise UserImportError("file_too_large")
        suffix = Path(filename).suffix.lower()
        if suffix == ".csv":
            rows = self._parse_csv(content)
        elif suffix == ".xlsx":
            rows = self._parse_xlsx(content)
        else:
            raise UserImportError("unsupported_file_type")
        required = {"username"} if domain_mode else {"username", "display_name", "domain"}
        if not rows:
            raise UserImportError("empty_import")
        headers = set(rows[0]["values"])
        if not required <= headers:
            raise UserImportError("missing_required_columns")
        if headers != required:
            raise UserImportError("unexpected_columns")
        data_rows = rows[1:]
        if len(data_rows) > MAX_IMPORT_ROWS:
            raise UserImportError("too_many_rows")
        parsed: list[dict[str, Any]] = []
        for item in data_rows:
            values = item["values"]
            if any(
                isinstance(value, str)
                and value.lstrip().startswith(("=", "+", "-", "@"))
                for value in values.values()
            ):
                raise UserImportError("formula_not_allowed")
            username = str(values.get("username") or "").strip()
            if not username:
                parsed.append({"row": item["row"], "username": "", "error": "username_required"})
                continue
            parsed.append({
                "row": item["row"],
                "username": username,
                "display_name": str(values.get("display_name") or "").strip() or None,
                "domain": str(values.get("domain") or "").strip() or None,
            })
        return parsed

    @staticmethod
    def _parse_csv(content: bytes) -> list[dict[str, Any]]:
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise UserImportError("csv_must_be_utf8") from exc
        reader = csv.DictReader(io.StringIO(text))
        headers = [str(name or "").strip() for name in (reader.fieldnames or [])]
        rows = [{"row": 1, "values": {name: name for name in headers}}]
        for number, raw in enumerate(reader, start=2):
            if number > MAX_IMPORT_ROWS + 1:
                raise UserImportError("too_many_rows")
            rows.append({
                "row": number,
                "values": {str(key or "").strip(): value for key, value in raw.items()},
            })
        return rows

    @staticmethod
    def _parse_xlsx(content: bytes) -> list[dict[str, Any]]:
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                members = archive.infolist()
                expanded = sum(member.file_size for member in members)
                compressed = sum(max(member.compress_size, 1) for member in members)
                if (
                    len(members) > MAX_XLSX_MEMBERS
                    or expanded > MAX_XLSX_EXPANDED_BYTES
                    or expanded > compressed * 100
                ):
                    raise UserImportError("xlsx_archive_limits_exceeded")
        except UserImportError:
            raise
        except zipfile.BadZipFile as exc:
            raise UserImportError("invalid_xlsx") from exc
        try:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
        except Exception as exc:  # openpyxl exposes several format-specific exceptions
            raise UserImportError("invalid_xlsx") from exc
        try:
            iterator = workbook.active.iter_rows(values_only=True)
            header_values = next(iterator, None)
            if header_values is None:
                return []
            headers = [str(value or "").strip() for value in header_values]
            rows = [{"row": 1, "values": {name: name for name in headers}}]
            for number, values in enumerate(iterator, start=2):
                if number > MAX_IMPORT_ROWS + 1:
                    raise UserImportError("too_many_rows")
                rows.append({
                    "row": number,
                    "values": {
                        header: value for header, value in zip(headers, values, strict=False)
                    },
                })
            return rows
        finally:
            workbook.close()

    async def _plan(
        self, *, rows: list[dict[str, Any]], domain: str | None,
    ) -> dict[str, Any]:
        new_by_username: dict[str, dict[str, Any]] = {}
        bindings: list[dict[str, str]] = []
        errors: list[dict[str, Any]] = []
        seen_bindings: set[tuple[str, str]] = set()
        usernames = sorted({str(row.get("username") or "") for row in rows if row.get("username")})
        existing_by_username = await self._db.get_users_by_usernames(usernames)
        for row in rows:
            username = row["username"]
            if row.get("error"):
                errors.append({
                    "row": row["row"], "username": username, "code": row["error"],
                })
                continue
            target_domain = domain or row.get("domain")
            try:
                resolve_domain(str(target_domain or ""))
            except Exception:
                errors.append({
                    "row": row["row"], "username": username, "code": "invalid_domain",
                })
                continue
            existing = existing_by_username.get(username)
            if existing is not None:
                if (
                    existing.get("status") != "active"
                    or existing.get("deleted_at") is not None
                    or existing.get("site_role") != "member"
                ):
                    errors.append({
                        "row": row["row"], "username": username, "code": "user_not_active",
                    })
                    continue
                key = (str(existing["id"]), str(target_domain))
                if key not in seen_bindings:
                    bindings.append({"user_id": key[0], "domain": key[1]})
                    seen_bindings.add(key)
                continue
            if domain is not None:
                errors.append({
                    "row": row["row"], "username": username, "code": "user_not_found",
                })
                continue
            planned = new_by_username.setdefault(username, {
                "username": username,
                "display_name": row.get("display_name"),
                "domains": [],
            })
            if planned["display_name"] != row.get("display_name"):
                errors.append({
                    "row": row["row"], "username": username,
                    "code": "conflicting_display_name",
                })
                continue
            if target_domain not in planned["domains"]:
                planned["domains"].append(target_domain)
        return {
            "new_users": list(new_by_username.values()),
            "bindings": bindings,
            "errors": errors,
            "total_rows": len(rows),
        }
