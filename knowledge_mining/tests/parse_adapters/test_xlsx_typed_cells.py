"""P1-7 回归：真实 xlsx parser 产类型化 cell 事实（Codex 审查）.

缺陷：``TableCell`` 契约有 ``value_type``/``normalized_value`` 但
native parser 不产——真实 Excel 新挖与历史回填后 cells 表两列恒 NULL，
date/千分位数值/类型聚合实际无法工作（A3 建立在空类型上）。

契约（从 parser 入口跑完整链路，不手工构造 TableCell）：
- number cell：value_type="number"，normalized_value=规范十进制文本
  （int 无小数点、float 最短表示）；
- date/datetime cell：value_type="date"，normalized_value=ISO
  （YYYY-MM-DD 或 YYYY-MM-DDTHH:MM:SS）；
- text cell：value_type="text"；
- 公式格：formula 保留 "=..."，value_type/normalized_value 按**计算结果**
  （display value）判定；
- 空/合并覆盖格：不产事实或 span 语义正确；
- 千分位显示格式：normalized_value 仍为规范数值（不受 display format
  污染——由 openpyxl number_format 不参与规范化保证）；
- 无法确认类型的值不强标（宁缺勿伪造）。
"""
from __future__ import annotations

import io
from datetime import datetime

import pytest

from knowledge_mining.mining.parse_adapters.native.native_xlsx import (
    NativeXlsxParser,
    XlsxNormalizer,
)

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
RAW_HASH = "ab" * 32


def _build_typed_xlsx() -> bytes:
    """单 sheet「台账」：覆盖 number/date/datetime/text/公式/空格/混合列."""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "台账"
    ws["A1"] = "设备"
    ws["B1"] = "功耗(W)"
    ws["C1"] = "投产日期"
    ws["D1"] = "巡检时间"
    ws["E1"] = "备注"
    ws["A2"] = "设备甲"
    ws["B2"] = 1234.5
    ws["C2"] = datetime(2026, 1, 5)
    ws["D2"] = datetime(2026, 9, 7, 8, 30, 0)
    ws["E2"] = "正常"
    ws["A3"] = "设备乙"
    ws["B3"] = 2000          # int
    ws["C3"] = datetime(2025, 12, 31)
    ws["D3"] = datetime(2026, 9, 8, 14, 0, 0)
    ws["E3"] = None          # 空格
    ws["A4"] = "合计"
    ws["B4"] = "=B2+B3"      # 公式（无缓存值场景由 display 兜底）
    ws["C4"] = None
    ws["D4"] = None
    ws["E4"] = "混合文本"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture
def typed_doc():
    parser = NativeXlsxParser()
    artifact = parser.parse(_build_typed_xlsx(), mime=XLSX_MIME)
    doc = XlsxNormalizer().normalize(artifact, source_raw_hash=RAW_HASH)
    return doc


def _cells_by_pos(doc) -> dict[tuple[int, int], object]:
    out: dict[tuple[int, int], object] = {}
    for asset in doc.structured_assets.values():
        for cell in asset.cells:
            out[(cell.row_index, cell.column_index)] = cell
    return out


# ---------------------------------------------------------------- number


def test_number_cell_typed_with_normalized_decimal(typed_doc):
    cells = _cells_by_pos(typed_doc)
    b2 = cells[(1, 1)]
    assert b2.value_type == "number"
    assert b2.normalized_value == "1234.5"
    assert b2.text == "1234.5"


def test_integer_number_normalized_without_decimal_point(typed_doc):
    cells = _cells_by_pos(typed_doc)
    b3 = cells[(2, 1)]
    assert b3.value_type == "number"
    assert b3.normalized_value == "2000"
    assert "." not in b3.normalized_value


# ---------------------------------------------------------------- date


def test_date_cell_iso_normalized(typed_doc):
    cells = _cells_by_pos(typed_doc)
    c2 = cells[(1, 2)]
    assert c2.value_type == "date"
    assert c2.normalized_value == "2026-01-05"


def test_datetime_cell_iso_with_time(typed_doc):
    cells = _cells_by_pos(typed_doc)
    d2 = cells[(1, 3)]
    assert d2.value_type == "date"
    assert d2.normalized_value == "2026-09-07T08:30:00"


# ---------------------------------------------------------------- text


def test_text_cell_typed_as_text(typed_doc):
    cells = _cells_by_pos(typed_doc)
    e2 = cells[(1, 4)]
    assert e2.value_type == "text"
    assert e2.normalized_value == "正常"


def test_header_cells_typed_as_text(typed_doc):
    cells = _cells_by_pos(typed_doc)
    for col in range(5):
        header = cells[(0, col)]
        assert header.value_type == "text", f"header col {col}"


# ---------------------------------------------------------------- 公式


def test_formula_kept_separate_from_display_value(typed_doc):
    cells = _cells_by_pos(typed_doc)
    b4 = cells[(3, 1)]
    assert b4.formula == "=B2+B3"
    # 公式格类型按计算结果（display）判定；无缓存值时不强标 number
    assert b4.value_type in (None, "number")


# ---------------------------------------------------------------- 空格/合并


def test_empty_cell_has_no_fabricated_type(typed_doc):
    """空格（None 值）不进 cells 事实——parser 只为有值格产事实（宁缺勿伪造）."""
    cells = _cells_by_pos(typed_doc)
    assert (2, 4) not in cells  # E3=None：无事实
    # 有值格全部有类型（text 列）
    assert cells[(3, 4)].value_type == "text"


def test_merged_origin_cell_keeps_span(typed_doc):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "M"
    ws["A1"] = "锚点"
    ws.merge_cells("A1:B2")
    buf = io.BytesIO()
    wb.save(buf)

    parser = NativeXlsxParser()
    artifact = parser.parse(buf.getvalue(), mime=XLSX_MIME)
    doc = XlsxNormalizer().normalize(artifact, source_raw_hash=RAW_HASH)
    cells = _cells_by_pos(doc)
    origin = cells[(0, 0)]
    assert origin.row_span == 2 and origin.column_span == 2
    assert origin.value_type == "text"


# ---------------------------------------------------------------- 全链投影


def test_extract_table_facts_consumes_typed_cells(typed_doc):
    """IR → TableFacts：类型化事实从 parser 端到 facts 投影可用."""
    from knowledge_mining.mining.table_assets.facts import extract_table_facts

    facts = extract_table_facts(typed_doc)
    assert facts.cells, "no cell facts extracted"
    # 抽任一 number 事实验证字段就位
    numbers = [f for f in facts.cells.values() if f.value_type == "number"]
    assert numbers, "no typed number facts"
    assert any(f.normalized_value == "1234.5" for f in numbers)
    dates = [f for f in facts.cells.values() if f.value_type == "date"]
    assert any(f.normalized_value == "2026-01-05" for f in dates)


def test_unconfirmed_values_not_force_typed(typed_doc):
    """混合列（文本+空+文本）不因猜测被强标 number/date."""
    cells = _cells_by_pos(typed_doc)
    col_e = [c for (r, c_), c in cells.items() if c_ == 4]
    for cell in col_e:
        assert cell.value_type in (None, "text")
